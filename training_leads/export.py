# -*- coding: utf-8 -*-
"""Выгрузка найденных служб обучения в Excel."""
from __future__ import annotations

import json
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [
    ("Организация", 34),
    ("ИНН", 12),
    ("КПП", 11),
    ("Регион", 22),
    ("Сайт организации", 26),
    ("Подразделение", 34),
    ("Должность", 34),
    ("ФИО", 26),
    ("E-mail", 28),
    ("Телефон", 20),
    ("Ссылка на страницу", 52),
    ("Источник", 15),
    ("Достоверность", 30),
    ("Фрагмент, где найдено", 60),
]


def rows(orgs: Dict[str, dict], leads: Dict[str, List[dict]]) -> List[list]:
    out = []
    for inn, org in orgs.items():
        name = org.get("name_short") or org.get("name_full") or ""
        found = leads.get(inn) or []
        if not found:
            # Компанию без находок из таблицы не выбрасываем: пустая строка
            # честнее молчания — видно, что искали и не нашли.
            out.append([
                name, inn, org.get("kpp", ""), org.get("region", ""),
                org.get("site", ""), "", "", "", "", "", "", "не найдено", "", "",
            ])
            continue
        for lead in found:
            out.append([
                name, inn, org.get("kpp", ""), org.get("region", ""),
                org.get("site", ""), lead.get("department", ""),
                lead.get("role", ""), lead.get("person", ""),
                lead.get("email", ""), ", ".join(lead.get("phones") or []),
                lead.get("url", ""), lead.get("source", ""),
                lead.get("grade", ""), lead.get("snippet", ""),
            ])
    return out


def to_xlsx(orgs, leads, path: str) -> int:
    wb = Workbook()
    ws = wb.active
    ws.title = "Обучение персонала"

    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="4F6228")
    ws.append([c[0] for c in COLUMNS])
    for i, (_, width) in enumerate(COLUMNS, start=1):
        ws.cell(row=1, column=i).font = head
        ws.cell(row=1, column=i).fill = fill
        ws.cell(row=1, column=i).alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width

    data = rows(orgs, leads)
    for row in data:
        ws.append(row)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(data) + 1}"
    for r in range(2, len(data) + 2):
        ws.cell(row=r, column=14).alignment = Alignment(wrap_text=False)
    wb.save(path)
    return len(data)
